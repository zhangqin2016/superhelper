"""Recalculate an XLSX into a separate verified copy, without user-profile macros."""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from lily_office_convert import convert


def recalculate(source, out_dir, timeout=60):
    source = Path(source).resolve()
    if source.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx is supported; preserve other workbook formats unchanged")
    if source.parent == Path(out_dir).resolve():
        raise ValueError("Use a separate output directory; the original must remain unchanged")
    original = load_workbook(source, data_only=False)
    try:
        formulas = {(sheet.title, cell.coordinate) for sheet in original
                    for row in sheet for cell in row if cell.data_type == "f"}
    finally:
        original.close()
    output = convert(str(source), str(out_dir), "xlsx:Calc MS Excel 2007 XML", timeout=timeout)
    values = load_workbook(output, data_only=True)
    expressions = load_workbook(output, data_only=False)
    errors = []
    try:
        for sheet, address in sorted(formulas):
            if sheet not in expressions or expressions[sheet][address].data_type != "f":
                errors.append(f"{sheet}!{address}: formula lost")
            elif values[sheet][address].value is None:
                errors.append(f"{sheet}!{address}: cache missing (empty result is not verified)")
        for sheet in values:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "e":
                        errors.append(f"{sheet.title}!{cell.coordinate}: {cell.value}")
    finally:
        values.close()
        expressions.close()
    if errors:
        raise ValueError("Recalculation not verified: " + "; ".join(errors[:30]))
    return {"status": "verified", "output": str(output), "formula_count": len(formulas),
            "scope": "formula presence, cached values and Excel errors; layout/features need separate review"}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source")
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--timeout", type=float, default=60)
    args = parser.parse_args()
    if not 0 < args.timeout <= 600:
        parser.error("timeout must be between 0 and 600 seconds")
    try:
        result = recalculate(args.source, args.out_dir, args.timeout)
    except Exception as error:
        print(json.dumps({"status": "unverified", "error": str(error)}))
        sys.exit(1)
    print(json.dumps(result))
