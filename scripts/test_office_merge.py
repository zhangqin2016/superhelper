import importlib.util
import io
import pathlib
import unittest
import zipfile
from docx import Document

script = pathlib.Path(__file__).resolve().parents[1] / 'resources/runtime-scripts/merge_office.py'
spec = importlib.util.spec_from_file_location('office_merge', script)
merge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(merge)


def package(document):
    stream = io.BytesIO()
    document.save(stream)
    with zipfile.ZipFile(stream, 'a') as archive:
        archive.writestr('custom/unknown.bin', b'untouched unknown part\x00\xff')
        archive.writestr('word/vbaProject.bin', b'opaque macro bytes\x00\xff')
    return stream.getvalue()


class OfficeMergeTest(unittest.TestCase):
    def setUp(self):
        doc = Document()
        doc.add_paragraph('original first')
        doc.add_paragraph('original second')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'left'
        table.cell(0, 1).text = 'right'
        doc.sections[0].header.paragraphs[0].text = 'Keep header'
        self.a = package(doc)

    def variant(self, update):
        doc = Document(io.BytesIO(self.a))
        update(doc)
        return package(doc)

    def test_disjoint_paragraphs_and_cells_preserve_package_parts(self):
        def local(doc):
            doc.paragraphs[0].text = 'private first'
            doc.tables[0].cell(0, 0).text = 'private left'
        def shared(doc):
            doc.paragraphs[1].text = 'shared second'
            doc.tables[0].cell(0, 1).text = 'shared right'
        result = merge.merge_docx(self.a, self.variant(local), self.variant(shared))
        doc = Document(io.BytesIO(result))
        self.assertEqual([p.text for p in doc.paragraphs], ['private first', 'shared second'])
        self.assertEqual([c.text for c in doc.tables[0].rows[0].cells], ['private left', 'shared right'])
        with zipfile.ZipFile(io.BytesIO(self.a)) as a, zipfile.ZipFile(io.BytesIO(result)) as output:
            for name in a.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(output.read(name), a.read(name), name)

    def test_same_paragraph_conflicts(self):
        def change(value):
            return self.variant(lambda doc: setattr(doc.paragraphs[0], 'text', value))
        with self.assertRaises(merge.MergeConflict):
            merge.merge_docx(self.a, change('private'), change('remote'))

    def test_structural_ambiguity_is_not_a_success(self):
        local = self.variant(lambda doc: doc.add_paragraph('new paragraph'))
        shared = self.variant(lambda doc: setattr(doc.paragraphs[0], 'text', 'remote'))
        with self.assertRaises(merge.MergeConflict):
            merge.merge_docx(self.a, local, shared)

    def test_duplicate_entries_refused(self):
        stream = io.BytesIO(self.a)
        with zipfile.ZipFile(stream, 'a') as archive:
            archive.writestr('word/document.xml', b'<bad/>')
        with self.assertRaises(merge.UnsupportedPackage):
            merge.merge_docx(stream.getvalue(), self.a, self.a)

    def replace_part(self, name, value):
        stream = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(self.a)) as source, zipfile.ZipFile(stream, 'w') as target:
            for part in source.namelist():
                target.writestr(part, value if part == name else source.read(part))
        return stream.getvalue()

    def test_opaque_part_conflict_preserves_both_inputs(self):
        with self.assertRaises(merge.MergeConflict):
            merge.merge_docx(self.a, self.replace_part('custom/unknown.bin', b'local'),
                             self.replace_part('custom/unknown.bin', b'remote'))

    def test_entities_and_path_traversal_are_refused(self):
        evil = b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><x>&secret;</x>'
        with self.assertRaises(merge.UnsupportedPackage):
            merge.merge_docx(self.replace_part('word/document.xml', evil), self.a, self.a)
        stream = io.BytesIO(self.a)
        with zipfile.ZipFile(stream, 'a') as archive:
            archive.writestr('../escape', b'bad')
        with self.assertRaises(merge.UnsupportedPackage):
            merge.merge_docx(stream.getvalue(), self.a, self.a)

    def test_relationship_edits_require_manual_resolution(self):
        with zipfile.ZipFile(io.BytesIO(self.a)) as archive:
            rels = archive.read('word/_rels/document.xml.rels')
        with self.assertRaises(merge.MergeConflict):
            merge.merge_docx(self.a, self.replace_part('word/_rels/document.xml.rels', rels + b'\n'), self.a)

    def test_signed_packages_and_member_limits_refused(self):
        for name, value in [('_xmlsignatures/sig1.xml', b'<signature/>'),
                            ('custom/large.bin', b'x' * (8 * 1024 * 1024 + 1))]:
            stream = io.BytesIO(self.a)
            with zipfile.ZipFile(stream, 'a', compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr(name, value)
            with self.assertRaises(merge.UnsupportedPackage):
                merge.merge_docx(stream.getvalue(), self.a, self.a)


if __name__ == '__main__':
    unittest.main()
