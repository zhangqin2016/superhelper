"""Native office regression: deep paths, cached formulas and preserved originals."""
import hashlib
from pathlib import Path
import sys
import tempfile
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "resources/runtime-scripts"))
from openpyxl import Workbook, load_workbook
import lily_office_convert as office
from lily_xlsx_recalc import recalculate


with tempfile.TemporaryDirectory(prefix="lily-recalc-test-") as root:
    root = Path(root)
    source = root / "formula.xlsx"
    book = Workbook()
    book.active.append([12, 3, "=A1-B1", "=SUM(A1:B1)"])
    book.save(source)
    before = hashlib.sha256(source.read_bytes()).digest()
    deep = root / ("nested-" * 9) / ("output-" * 7)
    deep.mkdir(parents=True)
    # Profile creation is independent of output depth and scoped to one call.
    with patch.object(office, "_run_with_profile", return_value=None) as run:
        office._run(str(source), str(deep), "xlsx", None, 25)
        profile = Path(run.call_args.args[-1])
        assert not profile.is_relative_to(deep)
        assert not profile.exists()
    result = recalculate(source, deep, 25)
    assert result["formula_count"] == 2
    values = load_workbook(result["output"], data_only=True)
    assert values.active["C1"].value == 9
    assert values.active["D1"].value == 15
    values.close()
    assert hashlib.sha256(source.read_bytes()).digest() == before
    try:
        recalculate(source, root)
        raise AssertionError("must not overwrite original")
    except ValueError:
        pass
    book.active["E1"] = "=1/0"
    bad = root / "error.xlsx"
    book.save(bad)
    try:
        recalculate(bad, deep, 25)
        raise AssertionError("Excel error must fail validation")
    except ValueError as error:
        assert "#DIV/0!" in str(error)
print("PASS: deep-path recalculation, cached values, original preservation, Excel errors")
